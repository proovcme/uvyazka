"""Literal readers; layout/OCR and engineering interpretation are separate stages."""
from __future__ import annotations

import csv
import io
from pathlib import Path
import zipfile
from importlib.metadata import version, PackageNotFoundError

READER_REVISION = "literal-readers-v1"
SUPPORTED = {".txt", ".md", ".csv", ".tsv", ".pdf", ".docx", ".xlsx", ".xlsm", ".xls"}
MAX_ROWS = 20000
MAX_CELLS = 300000
MAX_EXPANDED_BYTES = 64 * 1024 * 1024


def revision_for(suffix: str) -> str:
    package = {".pdf": "PyMuPDF", ".docx": "python-docx", ".xlsx": "openpyxl", ".xlsm": "openpyxl", ".xls":"xlrd"}.get(suffix)
    if not package:
        return READER_REVISION
    try:
        installed = version(package)
    except PackageNotFoundError:
        installed = "unavailable"
    # Excel row locators changed without invalidating unrelated PDF/Word evidence.
    adapter = ';row-locators-v2' if suffix in {'.xlsx','.xlsm'} else ''
    return READER_REVISION + ";" + package + "=" + installed + adapter


def read_blocks(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("Format is registered but requires another reader: " + suffix)
    if path.stat().st_size > MAX_EXPANDED_BYTES:
        raise ValueError("Source exceeds 64 MiB; use a bounded adapter")
    blocks = []

    def add(text: str, locator: dict) -> None:
        blocks.append({"text": text, "locator": locator})

    if suffix == ".pdf":
        try:
            import pymupdf
        except ImportError as exc:
            raise ValueError("Install uvyazka[readers] for PDF") from exc
        with pymupdf.open(path) as doc:
            if doc.needs_pass:
                raise ValueError("Encrypted PDF requires an approved readable copy")
            if len(doc) > 10000:
                raise ValueError("PDF exceeds 10000 pages; use a bounded adapter")
            for n, page in enumerate(doc, 1):
                add(page.get_text(), {"kind": "pdf_page", "pdf_page": n})
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("Install uvyazka[readers] for DOCX") from exc
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValueError("Expanded DOCX exceeds 64 MiB")
        doc = Document(path)
        for n, paragraph in enumerate(doc.paragraphs, 1):
            if paragraph.text.strip():
                add(paragraph.text, {"kind": "docx_paragraph", "paragraph": n, "note": "Not a physical page"})
        for table_no, table in enumerate(doc.tables, 1):
            for row_no, row in enumerate(table.rows, 1):
                add(" | ".join(cell.text for cell in row.cells), {"kind": "docx_table_row", "table": table_no, "row": row_no})
    elif suffix == '.xls':
        from .structured_readers import xls_units
        total, sheets=xls_units(path.read_bytes(),0,10000)
        count=0
        for sheet in sheets:
            if not sheet['tables']:raise ValueError('XLS sheet exceeds row/column/cell bounds')
            for row in sheet['tables'][0]['rows']:
                cells=[dict(cell=c['locator']['cell'],value=c['text'],type='xls_literal_or_cached',
                            formula_status='unknown') for c in row if c['text'] not in (None,'')]
                count+=len(cells)
                if count>MAX_CELLS or len(blocks)>=MAX_ROWS:raise ValueError('XLS exceeds total row/cell bounds')
                if cells:add('Лист: '+sheet['name']+'\n'+'\n'.join(c['cell']+': '+c['value'] for c in cells),
                             dict(kind='excel_row',sheet=sheet['name'],row=row[0]['row'],cells=cells,
                                  note='XLS formulas are not exposed; values may be saved calculation results'))
    elif suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("Install uvyazka[readers] for XLSX/XLSM") from exc
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValueError("Expanded workbook exceeds 64 MiB")
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = None
        try:
            cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
            count = 0
            for sheet in book:
                if sheet.max_row > MAX_ROWS or sheet.max_column > 512 or sheet.max_row * sheet.max_column > MAX_CELLS:
                    raise ValueError("Workbook sheet exceeds row/column/cell bounds")
                for row_number, (row, values) in enumerate(zip(sheet.iter_rows(), cached[sheet.title].iter_rows(values_only=True)),1):
                    cells = []
                    for cell, value in zip(row, values):
                        if cell.value is None:
                            continue
                        item = {"cell": cell.coordinate, "value": str(cell.value), "type": cell.data_type}
                        if cell.data_type == "f":
                            item.update(formula=str(cell.value), cached_value=str(value) if value is not None else None, needs_recalculation=value is None)
                        cells.append(item)
                    count += len(cells)
                    if count > MAX_CELLS or len(blocks) >= MAX_ROWS:
                        raise ValueError("Workbook exceeds total row/cell bounds")
                    if cells:
                        text = "Лист: " + sheet.title + "\n" + "\n".join(c["cell"] + ": " + c["value"] + (" [cached: " + str(c["cached_value"]) + "]" if "formula" in c else "") for c in cells)
                        add(text, {"kind": "excel_row", "sheet": sheet.title, "row": row_number, "cells": cells, "note": "Formulas/macros are not executed"})
        finally:
            book.close()
            if cached:
                cached.close()
    elif suffix in {".csv", ".tsv"}:
        text = path.read_text(encoding="utf-8-sig")
        delimiter = "\t" if suffix == ".tsv" else None
        dialect = None
        if delimiter is None:
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            except csv.Error:
                delimiter = ","
        rows = csv.reader(io.StringIO(text), dialect=dialect) if dialect else csv.reader(io.StringIO(text), delimiter=delimiter)
        count = 0
        for n, row in enumerate(rows, 1):
            count += len(row)
            if n > MAX_ROWS or len(row) > 512 or count > MAX_CELLS:
                raise ValueError("CSV exceeds row/column/cell bounds")
            if any(value != "" for value in row):
                add(" | ".join(f"C{i}: {value}" for i, value in enumerate(row, 1)), {"kind": "csv_row", "row": n, "cells": row})
    else:
        text = path.read_text(encoding="utf-8-sig")
        for n, line in enumerate(text.splitlines(), 1):
            if line.strip():
                add(line, {"kind": "text_line", "line": n})
    return blocks
