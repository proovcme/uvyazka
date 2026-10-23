"""Literal grids and source units. No calculation, engineering approval or page guessing."""
from io import BytesIO
from importlib.metadata import version, PackageNotFoundError
from functools import lru_cache
import csv
import zipfile
import xml.etree.ElementTree as ET
import warnings

FORMATS = {'.pdf', '.xlsx', '.xlsm', '.xls', '.docx', '.csv', '.tsv'}
REVISION = 'source-grids-v3'
MAX_BYTES = 64 * 1024 * 1024
MAX_CELLS = 300000
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


@lru_cache(maxsize=32)
def engine(suffix):
    package = {'.pdf': 'PyMuPDF', '.xlsx': 'openpyxl', '.xlsm': 'openpyxl', '.xls':'xlrd'}.get(suffix)
    try:
        dependency = version(package) if package else 'stdlib'
    except PackageNotFoundError:
        dependency = 'unavailable'
    return f'{REVISION}/{package or "xml-csv"}:{dependency}'


def archive(data):
    z = zipfile.ZipFile(BytesIO(data))
    if len(z.infolist()) > 10000 or sum(i.file_size for i in z.infolist()) > MAX_BYTES:
        z.close()
        raise ValueError('Expanded document exceeds 64 MiB or 10000 entries')
    if any(i.flag_bits & 1 for i in z.infolist()):
        z.close()
        raise ValueError('Encrypted document is unsupported')
    return z


def cell(row, column, text, locator, **extra):
    return dict(row=row, column=column, text=text, state='empty' if text == '' else 'value',
                locator=locator, **extra)


def table(name, rows, columns, **extra):
    if len(rows) > 20000 or columns > 512 or sum(len(r) for r in rows) > MAX_CELLS:
        raise ValueError('Grid exceeds 20000 rows, 512 columns or 300000 cells')
    return dict(name=name, row_count=len(rows), column_count=columns, rows=rows, **extra)


def excel_units(data, start, limit):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, range_boundaries
    # Read-only interpretation; neither VBA, links nor formula evaluation is run.
    with archive(data) as z:
        # Normal-mode openpyxl retains cells for merges/styles. Bound object allocation before loading.
        count=0
        for name in z.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                with z.open(name) as stream:
                    for event,node in ET.iterparse(stream,events=('end',)):
                        if node.tag==S+'c':
                            count+=1
                            if count>MAX_CELLS:raise ValueError('Workbook exceeds 300000 serialized cells')
                        node.clear()
        wb_xml=ET.fromstring(z.read('xl/workbook.xml'))
        definitions=[dict(n.attrib,text=n.text) for n in wb_xml.iter(S+'definedName')]
        with warnings.catch_warnings(record=True) as diagnostics:
            warnings.simplefilter('always')
            book = load_workbook(BytesIO(data), data_only=False, keep_links=False)
            cached = load_workbook(BytesIO(data), data_only=True, keep_links=False)
        try:
            result = []
            for index in range(start, min(len(book.worksheets), start + limit)):
                sheet, values = book.worksheets[index], cached.worksheets[index]
                # Exact XML presence distinguishes absent cells from explicit empty cells.
                xml = ET.fromstring(z.read(sheet.path.lstrip('/'))) if sheet.path.lstrip('/') in z.namelist() else None
                # openpyxl assigns worksheet paths only during save; use relationship order fallback.
                if xml is None:
                    wb = ET.fromstring(z.read('xl/workbook.xml'))
                    rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
                    rel_id = list(wb.find(S+'sheets'))[index].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    target = next(r.get('Target') for r in rels if r.get('Id') == rel_id)
                    import posixpath
                    target = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
                    xml = ET.fromstring(z.read(target))
                raw = {c.get('r'): c for c in xml.iter(S+'c')}
                max_row = max((c.row for c in sheet._cells.values()), default=0)
                max_col = max((c.column for c in sheet._cells.values()), default=0)
                merges = [str(r) for r in sheet.merged_cells.ranges]
                unit = dict(kind='excel_sheet', name=sheet.title, index=index+1,
                            physical_page=False, warnings=['reader_warning:'+str(w.message)[:400] for w in diagnostics],
                            tables=[], print_area=str(sheet.print_area),
                            raw_print_definitions=[d for d in definitions if d.get('localSheetId')==str(index)
                                                   and d.get('name') in {'_xlnm.Print_Area','_xlnm.Print_Titles'}],
                            print_title_rows=sheet.print_title_rows, print_title_cols=sheet.print_title_cols,
                            row_breaks=[b.id for b in sheet.row_breaks.brk],
                            column_breaks=[b.id for b in sheet.col_breaks.brk], state=sheet.sheet_state,
                            merged_ranges=merges,
                            native_tables=[dict(name=t.name, ref=t.ref, header_rows=t.headerRowCount,
                                                totals_rows=t.totalsRowCount) for t in sheet.tables.values()])
                if max_row > 20000 or max_col > 512 or max_row * max_col > MAX_CELLS:
                    unit['warnings'].append('grid_exceeds_budget_not_extracted')
                    result.append(unit)
                    continue
                owners = {}
                for ref in merges:
                    c1, r1, c2, r2 = range_boundaries(ref)
                    for r in range(r1, r2+1):
                        for c in range(c1, c2+1):
                            owners[(r, c)] = dict(anchor=f'{get_column_letter(c1)}{r1}',
                                                 row_span=r2-r1+1, column_span=c2-c1+1)
                rows = []
                for r in range(1, max_row+1):
                    output = []
                    for c in range(1, max_col+1):
                        address = f'{get_column_letter(c)}{r}'
                        source = raw.get(address)
                        value = sheet.cell(r,c).value
                        formula = value if sheet.cell(r,c).data_type == 'f' else None
                        cached_value = values.cell(r,c).value
                        text = '' if value is None else str(value)
                        item = cell(r,c,text, dict(kind='excel_cell',sheet=sheet.title,cell=address),
                                    value_type=sheet.cell(r,c).data_type, formula=formula,
                                    cached_value=None if cached_value is None else str(cached_value),
                                    number_format=sheet.cell(r,c).number_format,
                                    hidden_row=bool(sheet.row_dimensions[r].hidden),
                                    hidden_column=bool(sheet.column_dimensions[get_column_letter(c)].hidden))
                        if source is None:
                            item['state'] = 'absent'
                        else:
                            v = source.find(S+'v')
                            item['raw_xml_value'] = v.text if v is not None else None
                            if source.get('t') == 'n' and not formula and v is not None and v.text is not None:
                                item['text'] = v.text  # Preserve exact numeric serialization.
                            if formula:
                                item['cache_status'] = 'present' if v is not None and v.text is not None else 'missing'
                                item['state'] = 'formula'
                        if (r,c) in owners:
                            item['merge'] = owners[(r,c)]
                            if address != item['merge']['anchor']:
                                item.update(state='merged_continuation', text=None)
                        output.append(item)
                    rows.append(output)
                unit['tables'] = [table(sheet.title,rows,max_col, container_kind='worksheet_grid',
                                        header_status='not_inferred', merged_ranges=merges)]
                result.append(unit)
            return len(book.worksheets), result
        finally:
            book.close()
            cached.close()


def word_units(data, start, limit):
    with archive(data) as z:
        parts = sorted(n for n in z.namelist() if n == 'word/document.xml' or
                       (n.startswith(('word/header','word/footer')) and n.endswith('.xml')))
        parts.sort(key=lambda n: (n != 'word/document.xml', n))
        output = []
        for index in range(start, min(len(parts),start+limit)):
            part = parts[index]
            root = ET.fromstring(z.read(part))
            paths = {}
            def walk(node, path):
                paths[id(node)] = path
                counts={}
                for child in node:
                    counts[child.tag]=counts.get(child.tag,0)+1
                    walk(child, path+'/'+child.tag.split('}')[-1]+f'[{counts[child.tag]}]')
            walk(root,'/'+root.tag.split('}')[-1])
            unit = dict(index=index+1,kind='word_part',name=part,physical_page=False,
                        warnings=[],tables=[],blocks=[])
            all_tables = list(root.iter(W+'tbl'))
            parents = {child: parent for parent in root.iter() for child in parent}
            for node in root.iter():
                if node.tag in {W+'p',W+'tbl'}:
                    unit['blocks'].append(dict(kind=node.tag.split('}')[-1],xml_path=paths[id(node)]))
            # Store only navigation context, not a second unbounded copy of all paragraph text.
            unit['block_count'] = len(unit['blocks'])
            unit['blocks'] = unit['blocks'][:100]
            for ordinal, t in enumerate(all_tables,1):
                rows = []; active = {}; header_rows = []
                ancestors = []; p = parents.get(t)
                while p is not None:
                    if p.tag == W+'tbl': ancestors.append(all_tables.index(p)+1)
                    p = parents.get(p)
                for rn, tr in enumerate(t.findall(W+'tr'),1):
                    row = []; before = tr.find(W+'trPr/'+W+'gridBefore')
                    col = int(before.get(W+'val','0'))+1 if before is not None else 1
                    if col<1 or col>513:raise ValueError('Invalid DOCX gridBefore')
                    for omitted in range(1,col):
                        missing=cell(rn,omitted,None,dict(kind='word_omitted_cell',part=part,table=ordinal,
                            xml_path=paths[id(tr)],grid_omission='before'))
                        missing['state']='absent';row.append(missing)
                    repeated=tr.find(W+'trPr/'+W+'tblHeader')
                    if repeated is not None and repeated.get(W+'val','true') not in {'false','0','off'}:header_rows.append(rn)
                    next_active = {}
                    for tc in tr.findall(W+'tc'):
                        span = tc.find(W+'tcPr/'+W+'gridSpan')
                        width = int(span.get(W+'val','1')) if span is not None else 1
                        if width < 1 or col+width > 513: raise ValueError('Invalid DOCX table span')
                        # Nested tables have their own catalog entries, not duplicated parent quantities.
                        def own_text(n):
                            if n.tag == W+'tbl': return ''
                            if n.tag == W+'t': return n.text or ''
                            if n.tag in {W+'br',W+'tab'}: return '\n' if n.tag == W+'br' else '\t'
                            return ''.join(own_text(child) + ('\n' if child.tag == W+'p' else '') for child in n)
                        text = own_text(tc).rstrip('\n')
                        locator = dict(kind='word_cell',part=part,table=ordinal,xml_path=paths[id(tc)])
                        item = cell(rn,col,text,locator,column_span=width,row_span=1)
                        vertical = tc.find(W+'tcPr/'+W+'vMerge')
                        if vertical is not None:
                            if vertical.get(W+'val') == 'restart':
                                next_active[(col,width)] = item
                            elif (col,width) in active:
                                anchor = active[(col,width)]
                                anchor['row_span'] += 1
                                next_active[(col,width)] = anchor
                                item.update(state='merged_continuation',text=None,
                                            merge_owner=dict(row=anchor['row'],column=anchor['column']))
                            else:
                                unit['warnings'].append(f'orphan_vertical_merge:{ordinal}:{rn}:{col}')
                                item['state'] = 'unresolved_merge'
                        row.append(item)
                        for cc in range(col+1,col+width):
                            row.append(cell(rn,cc,None,locator))
                            row[-1].update(state='merged_continuation',merge_owner=dict(row=rn,column=col))
                        col += width
                    after=tr.find(W+'trPr/'+W+'gridAfter')
                    trailing=int(after.get(W+'val','0')) if after is not None else 0
                    if trailing<0 or col+trailing>513:raise ValueError('Invalid DOCX gridAfter')
                    for omitted in range(col,col+trailing):
                        missing=cell(rn,omitted,None,dict(kind='word_omitted_cell',part=part,table=ordinal,
                            xml_path=paths[id(tr)],grid_omission='after'))
                        missing['state']='absent';row.append(missing)
                    rows.append(row); active = next_active
                grid=t.find(W+'tblGrid')
                columns = max(len(grid) if grid is not None else 0,max((c['column'] for row in rows for c in row),default=0))
                unit['tables'].append(table(f'Таблица {ordinal}',rows,columns,
                    source_table=ordinal,parent_tables=ancestors,xml_path=paths[id(t)],
                    header_rows=header_rows,header_status='native_repeat_flag_only'))
            output.append(unit)
        return len(parts),output


def xls_units(data,start,limit):
    import xlrd
    book=xlrd.open_workbook(file_contents=data,on_demand=True,formatting_info=True)
    try:
        units=[]
        for index in range(start,min(book.nsheets,start+limit)):
            s=book.sheet_by_index(index)
            unit=dict(index=index+1,kind='excel_sheet',name=s.name,physical_page=False,
                      warnings=['xls_formula_expressions_unavailable_values_may_be_cached_results'],
                      tables=[],date_system=book.datemode)
            if s.nrows>20000 or s.ncols>512 or s.nrows*s.ncols>MAX_CELLS:
                unit['warnings'].append('grid_exceeds_budget_not_extracted');units.append(unit);continue
            merges={}
            for r1,r2,c1,c2 in s.merged_cells:
                anchor=xlrd.formula.colname(c1)+str(r1+1)
                for r in range(r1,r2):
                    for c in range(c1,c2):merges[r,c]=dict(anchor=anchor,row_span=r2-r1,column_span=c2-c1)
            rows=[]
            for r in range(s.nrows):
                row=[]
                for c in range(s.ncols):
                    value=s.cell(r,c);address=xlrd.formula.colname(c)+str(r+1)
                    raw=str(bool(value.value)) if value.ctype==xlrd.XL_CELL_BOOLEAN else str(value.value)
                    item=cell(r+1,c+1,raw,dict(kind='excel_cell',sheet=s.name,cell=address),
                              value_type=value.ctype,formula_status='unknown_not_exposed_by_xlrd')
                    if value.ctype==xlrd.XL_CELL_EMPTY:item.update(state='unresolved_empty',text=None)
                    if value.ctype==xlrd.XL_CELL_BLANK:item.update(state='empty',text='')
                    if value.ctype in {xlrd.XL_CELL_NUMBER,xlrd.XL_CELL_DATE}:
                        item['native_float_hex']=value.value.hex()
                    if value.xf_index is not None:
                        xf=book.xf_list[value.xf_index];fmt=book.format_map.get(xf.format_key)
                        item['number_format']=fmt.format_str if fmt else None
                    if (r,c) in merges:
                        item['merge']=merges[r,c]
                        if item['merge']['anchor']!=address:item.update(state='merged_continuation',text=None)
                    row.append(item)
                rows.append(row)
            unit['tables']=[table(s.name,rows,s.ncols,container_kind='worksheet_grid',header_status='not_inferred')]
            units.append(unit)
        return book.nsheets,units
    finally:book.release_resources()


def pdf_units(pdf, start, limit):
    from .document_tables import detect, table_matrix, navigation
    from .measurements import COORDINATES
    output = []
    for index in range(start,min(len(pdf),start+limit)):
        page = pdf[index]; rotation = page.rotation; page.set_rotation(0)
        text = page.get_text('text')
        unit = dict(index=index+1,kind='pdf_page',name=f'Страница PDF {index+1}',physical_page=True,
                    original_rotation=rotation,bounds=list(page.rect),crop_box=list(page.cropbox),
                    coordinate_system=COORDINATES,warnings=[],tables=[],
                    navigation_candidates=navigation(text[:1000000],index+1),
                    printed_sheet_number=None,views_status='requires_identification_and_separate_scales')
        if not text.strip(): unit['warnings'].append('no_text_layer_OCR_or_visual_read_required')
        try:
            for ordinal,t in enumerate(detect(page).tables,1):
                matrix = table_matrix(t); rows = []
                for rn, values in enumerate(matrix,1):
                    row = []
                    for cn,value in enumerate(values,1):
                        bounds = t.rows[rn-1].cells[cn-1]
                        item = cell(rn,cn,value,dict(kind='pdf_cell',pdf_page=index+1,
                            table=ordinal,bbox=list(bounds) if bounds else None,coordinate_system=COORDINATES))
                        if value is None: item['state'] = 'unresolved_or_merged'
                        row.append(item)
                    rows.append(row)
                unit['tables'].append(table(f'Таблица {ordinal}',rows,t.col_count,bbox=list(t.bbox),
                    header_status='not_inferred',coordinate_system=COORDINATES))
        except Exception as exc:
            unit['warnings'].append('table_detection:'+type(exc).__name__+':'+str(exc)[:200])
        output.append(unit)
    return len(pdf),output


def extract(data, suffix, start=0, limit=5, pdf=None):
    if suffix == '.pdf': return pdf_units(pdf,start,limit)
    if suffix in {'.xlsx','.xlsm'}: return excel_units(data,start,limit)
    if suffix=='.xls':return xls_units(data,start,limit)
    if suffix == '.docx': return word_units(data,start,limit)
    if suffix in {'.csv','.tsv'}:
        if start: return 1,[]
        try: text = data.decode('utf-8-sig')
        except UnicodeDecodeError: text = data.decode('cp1251')
        rows=[]; count=0
        try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t') if suffix=='.csv' else None
        except csv.Error:dialect=None
        reader=csv.reader(text.splitlines(keepends=True),dialect=dialect) if dialect else csv.reader(text.splitlines(keepends=True),delimiter='\t' if suffix=='.tsv' else ',')
        for r,values in enumerate(reader,1):
            count+=len(values)
            if r>20000 or len(values)>512 or count>MAX_CELLS:raise ValueError('Delimited grid exceeds budget')
            rows.append([cell(r,c,v,dict(kind='delimited_cell',row=r,column=c)) for c,v in enumerate(values,1)])
        return 1,[dict(index=1,kind='delimited_data',name='Данные',physical_page=False,warnings=[],
                       tables=[table('Данные',rows,max(map(len,rows),default=0),header_status='not_inferred')])]
    raise ValueError('Unsupported structured format: '+suffix)
