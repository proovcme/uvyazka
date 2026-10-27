import copy
import tempfile
import io
import time
import zipfile
import unittest
from pathlib import Path

from uvyazka.construction.geometry import box_model
from uvyazka.construction.clashes import run_clashes
from uvyazka.construction.clash_geometry import Mesh, Budget, compare
from uvyazka.construction.coordination_store import CoordinationStore
from uvyazka.construction.viewer_workspace import RevisionConflict
from uvyazka.core import Workspace
from uvyazka.construction.clash_export import bcf,xlsx,camera_world
from uvyazka.construction.clash_jobs import ClashJobs,compare_runs


def fixture(boxes):
    model,binary=box_model('Analytical synthetic solids',[
        {'id':str(i),'name':'Synthetic solid '+str(i),'bbox_min_m':a,'bbox_max_m':b,'geometry_basis':'Exact synthetic cuboid'} for i,(a,b) in enumerate(boxes)],[],[])
    for e in model['elements']:e['mesh']['source']['method']='analytical_synthetic_solid'
    return model,binary


def rule(kind='hard',**extra):
    return {'version':1,'id':'synthetic-test','name':'Synthetic intersection','kind':kind,'left':{'ids':['0']},'right':{'ids':['1']},**extra}


class ClashGeometryTests(unittest.TestCase):
    def run_boxes(self,boxes,**options):
        m,b=fixture(boxes);return run_clashes(m,b,rule(**options))

    def test_intersection_containment_and_thin_crossing(self):
        for boxes in [
            [([0,0,0],[2,2,2]),([1,1,1],[3,3,3])],
            [([0,0,0],[4,4,4]),([1,1,1],[2,2,2])],
            [([-5,-.1,-.1],[5,.1,.1]),([-.1,-5,-.1],[.1,5,.1])],
            [([0,0,0],[2,2,2]),([0,0,0],[2,2,2])],
            [([0,0,0],[.3,5,3]),([.05,.5,.5],[.25,1,1])],
        ]:
            with self.subTest(boxes=boxes):
                r=self.run_boxes(boxes)
                self.assertTrue(r['complete']);self.assertEqual(len(r['issues']),1)
                self.assertEqual(r['issues'][0]['kind'],'intersection')
                self.assertEqual(r['issues'][0]['evidence']['proof'],'interior_witness')

    def test_touch_is_separate_from_penetration_and_gap_is_not_collision(self):
        r=self.run_boxes([([0,0,0],[1,1,1]),([1,0,0],[2,1,1])])
        self.assertTrue(r['complete']);self.assertFalse(r['issues']);self.assertEqual(len(r['contacts']),1)
        r=self.run_boxes([([0,0,0],[1,1,1]),([1.2,0,0],[2.2,1,1])])
        self.assertFalse(r['issues']);self.assertEqual(r['coverage']['checked_pairs'],0)

    def test_no_claim_of_above_tolerance_penetration_for_a_thin_overlap(self):
        r=self.run_boxes([([0,0,0],[1,1,1]),([.9999,0,0],[2,1,1])],tolerance_mm=1)
        self.assertTrue(all(i['kind']!='intersection' for i in r['issues']))

    def test_clearance_and_units(self):
        r=self.run_boxes([([0,0,0],[1,1,1]),([1.05,0,0],[2.05,1,1])],kind='clearance',clearance_mm=100)
        self.assertAlmostEqual(r['issues'][0]['evidence']['distance_m'],.05,places=5)
        self.assertEqual(r['issues'][0]['kind'],'clearance')
        r=self.run_boxes([([0,0,0],[1,1,1]),([1.05,0,0],[2.05,1,1])],kind='clearance',clearance_mm=40)
        self.assertFalse(r['issues'])

    def test_source_feet_and_millimetres_do_not_rescale_already_normalized_meshes(self):
        for scale in (.3048,.001):
            m,b=fixture([([0,0,0],[1,1,1]),([1.05,0,0],[2.05,1,1])])
            m['coordinate_system']['length_to_m']=scale
            r=run_clashes(m,b,rule('clearance',clearance_mm=100))
            self.assertAlmostEqual(r['issues'][0]['evidence']['distance_m'],.05,places=5)
        for scale in (None,0,float('nan')):
            m['coordinate_system']['length_to_m']=scale
            with self.assertRaises(ValueError):run_clashes(m,b,rule())

    def test_sparse_model_uses_spatial_candidates_instead_of_a_quadratic_scan(self):
        m,b=fixture([([i*3,0,0],[i*3+1,1,1]) for i in range(500)])
        r=run_clashes(m,b,{**rule(),'left':{},'right':{}})
        self.assertTrue(r['complete']);self.assertFalse(r['issues'])
        self.assertEqual(r['coverage']['checked_pairs'],0)
        self.assertLess(r['coverage']['broadphase_tests'],10000)

    def test_unread_and_approximate_cannot_be_complete(self):
        m,b=fixture([([0,0,0],[2,2,2]),([1,1,1],[3,3,3])])
        m['elements'][0]['mesh']['source']['method']='document_bounding_box'
        r=run_clashes(m,b,rule());self.assertFalse(r['complete']);self.assertEqual(r['issues'][0]['kind'],'question')
        m['elements'][0]['mesh']=None
        r=run_clashes(m,b,rule());self.assertFalse(r['complete']);self.assertEqual(len(r['coverage']['unread']),1)
        with self.assertRaises(ValueError):run_clashes(m,b[:-1],rule())

    def test_empty_selection_budget_cancel_and_revision_pinned_exclusion(self):
        m,b=fixture([([0,0,0],[2,2,2]),([1,1,1],[3,3,3])])
        r=run_clashes(m,b,{**rule(),'left':{'ids':['absent']}});self.assertFalse(r['complete'])
        r=run_clashes(m,b,rule(max_triangle_tests=1));self.assertEqual(r['status'],'partial');self.assertFalse(r['complete'])
        r=run_clashes(m,b,rule(),cancelled=lambda:True);self.assertEqual(r['status'],'cancelled')
        exclusion={'a':'0','b':'1','reason':'Synthetic approved exception','source_sha256':m['source']['sha256']}
        r=run_clashes(m,b,rule(excluded_pairs=[exclusion]));self.assertEqual(r['coverage']['excluded_pairs'],1)
        exclusion['source_sha256']='f'*64
        self.assertEqual(len(run_clashes(m,b,rule(excluded_pairs=[exclusion]))['issues']),1)

    def test_pairs_are_deduplicated_and_service_zones_have_basis(self):
        m,b=fixture([([0,0,0],[1,1,1]),([1.5,0,0],[2.5,1,1])])
        r=run_clashes(m,b,{**rule('service'),'zones':[{'id':'door','owner_id':'0','bbox_min_m':[1,0,0],'bbox_max_m':[2,1,1],'basis':'Synthetic removal path'}]})
        self.assertEqual(r['issues'][0]['a']['id'],'0');self.assertEqual(r['issues'][0]['kind'],'intersection')
        r=run_clashes(m,b,{**rule('clearance',clearance_mm=1000),'left':{},'right':{}})
        self.assertEqual(r['coverage']['checked_pairs'],1)


class CoordinationStoreTests(unittest.TestCase):
    def test_shared_version_history_idempotency_and_scene_integrity(self):
        import hashlib
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);(root/'source').mkdir();w=Workspace(root/'state');w.create_project('p','Synthetic',root/'source')
            store=CoordinationStore(w);saved=store.save('p','clash_template','t',rule(),0,'request')
            self.assertEqual(store.save('p','clash_template','t',rule(),0,'request'),saved)
            with self.assertRaises(RevisionConflict):store.save('p','clash_template','t',rule(),0,'another')
            with self.assertRaises(ValueError):store.save('p','clash_template','t',{**rule(),'name':'Other'},0,'request')
            self.assertEqual(store.get('p','clash_template','t')['payload'],rule())
            m,b=fixture([([0,0,0],[1,1,1])]);scene=store.begin_scene('p',m,hashlib.sha256(b).hexdigest())
            self.assertFalse(scene['ready'])
            with self.assertRaises(ValueError):store.scene('p',scene['scene_id'])
            self.assertTrue(store.scene_chunk('p',scene['scene_id'],0,b)['ready'])
            self.assertTrue(store.scene_chunk('p',scene['scene_id'],0,b)['ready'])
            self.assertEqual(store.scene('p',scene['scene_id']),(m,b))
            (store.scene_folder('p',scene['scene_id'])/'geometry.bin').write_bytes(b'changed')
            with self.assertRaises(ValueError):store.scene('p',scene['scene_id'])

    def test_job_persists_and_mcp_client_can_read_result(self):
        import hashlib
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);(root/'source').mkdir();w=Workspace(root/'state');w.create_project('p','Synthetic',root/'source')
            m,b=fixture([([0,0,0],[2,2,2]),([1,1,1],[3,3,3])]);store=CoordinationStore(w)
            scene=store.begin_scene('p',m,hashlib.sha256(b).hexdigest());store.scene_chunk('p',scene['scene_id'],0,b)
            jobs=ClashJobs(w);job=jobs.start('p',scene['scene_id'],rule(),'native-job')
            deadline=time.monotonic()+10
            while job['status'] in {'queued','running'} and time.monotonic()<deadline:
                time.sleep(.02);job=jobs.get('p',job['job_id'])
            self.assertEqual(job['status'],'complete',job)
            self.assertEqual(jobs.start('p',scene['scene_id'],rule(),'native-job')['job_id'],job['job_id'])
            self.assertEqual(store.get('p','clash_run',job['job_id'])['payload']['issues'],job['report']['issues'])


class ClashExportTests(unittest.TestCase):
    def report(self):
        m,b=fixture([([0,0,0],[2,2,2]),([1,1,1],[3,3,3])])
        m['elements'][0]['name']='=1+1';m['elements'][0]['revit']={'element_id':'123'}
        return run_clashes(m,b,rule())

    def test_bcf_schemas_and_no_fabricated_ifc_guid(self):
        from xml.etree import ElementTree as E
        r=self.report()
        for version in ('2.1','3.0'):
            with zipfile.ZipFile(io.BytesIO(bcf(r,version=version))) as z:
                for name in z.namelist():
                    if name.endswith('.bcfv'):
                        info=E.fromstring(z.read(name));c=info.find('.//Component')
                        self.assertNotIn('IfcGuid',c.attrib);self.assertEqual(c.findtext('AuthoringToolId'),'123')
                    if name.endswith(('.bcfv','.bcf')):
                        schemas=Path(__file__).resolve().parents[1]/'.uvyazka-build/bcf-schemas'/version
                        if schemas.is_dir():
                            from lxml import etree
                            schema=etree.XMLSchema(etree.parse(str(schemas/('visinfo.xsd' if name.endswith('.bcfv') else 'markup.xsd'))))
                            schema.assertValid(etree.fromstring(z.read(name)))

    def test_xlsx_numbers_and_literal_formula_like_text(self):
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(xlsx(self.report())))
        self.assertEqual(book.sheetnames,['Замечания','Покрытие','Источники'])
        self.assertEqual(book['Замечания']['F2'].value,'=1+1');self.assertEqual(book['Замечания']['F2'].data_type,'s')
        self.assertEqual(book['Замечания']['L2'].value,0)

    def test_comparison_will_not_resolve_missing_coverage_or_unproven_lineage(self):
        r=self.report();after=copy.deepcopy(r);after['issues']=[];after['complete']=False
        self.assertEqual(compare_runs(r,after)['items'][0]['status'],'question')
        after['complete']=True;self.assertEqual(compare_runs(r,after)['items'][0]['status'],'not_detected_requires_review')
        after['source_sha256']='e'*64;self.assertFalse(compare_runs(r,after)['comparable'])


if __name__=='__main__':unittest.main()
