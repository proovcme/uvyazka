"""Synthetic literal grids, independent readings and shared HTTP/MCP persistence."""
import json
from pathlib import Path
import tempfile
import unittest
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from http.server import ThreadingHTTPServer
from threading import Thread

from uvyazka.core import Workspace,connection
from uvyazka.ingest import scan
from uvyazka.project_structure import ProjectStructure


class StructureTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.sources=self.root/'input';self.sources.mkdir()
        self.workspace=Workspace(self.root/'state');self.workspace.create_project('a','Synthetic',self.sources)

    def tearDown(self):self.temp.cleanup()

    def prepare(self):
        scan(self.workspace,'a');self.store=ProjectStructure(self.workspace,'a')
        return self.store

    def test_excel_merges_presence_exact_values_formula_cache_and_resume(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.table import Table
        book=Workbook();s=book.active;s.title='Оборудование'
        s.append(['Наименование','Количество',None]);s.merge_cells('B1:C1')
        s.append(['Насос',0,False]);s.append(['',None,'=B2+1'])
        s['C4']='000123';s['B4']='';s.row_dimensions[2].hidden=True
        s.print_area='A1:C4';s.print_title_rows='1:1'
        s=book.create_sheet('Данные');s.append(['Марка','Количество']);s.append(['П1',10]);s.add_table(Table(displayName='Equipment',ref='A1:B2'))
        book.save(self.sources/'data.xlsx')
        store=self.prepare();first=store.build(unit_limit=1)
        self.assertEqual(first['coverage']['pending_documents'],1)
        self.assertEqual(first['coverage']['indexed_units'],1)
        store=ProjectStructure(self.workspace,'a');last=store.build(unit_limit=1)
        self.assertEqual(last['coverage']['pending_documents'],0)
        self.assertEqual(last['coverage']['tables'],2)
        sheets=store.catalog()['items'];self.assertTrue(all(not s['physical_page'] for s in sheets))
        sheet=next(s for s in sheets if s['name']=='Оборудование')
        self.assertEqual(sheet['merged_ranges'],['B1:C1']);self.assertIn('A',sheet['print_area'])
        data=store.read(sheet['table_ids'][0]);cells={(c['row'],c['column']):c for row in data['rows'] for c in row}
        self.assertEqual(cells[1,3]['state'],'merged_continuation');self.assertIsNone(cells[1,3]['text'])
        self.assertEqual(cells[1,2]['merge']['column_span'],2)
        self.assertEqual(cells[2,2]['text'],'0');self.assertEqual(cells[2,3]['text'],'False')
        self.assertEqual(cells[3,2]['state'],'absent');self.assertEqual(cells[4,2]['state'],'empty')
        self.assertEqual(cells[4,3]['text'],'000123')
        self.assertEqual(cells[3,3]['formula'],'=B2+1');self.assertIsNone(cells[3,3]['cached_value'])
        self.assertEqual(cells[3,3]['cache_status'],'missing')
        self.assertTrue(cells[2,1]['hidden_row']);self.assertTrue(data['current_registered_version'])
        self.assertEqual(store.build()['processed'],[])

    def test_word_merges_nested_headers_order_and_no_fake_pages(self):
        from docx import Document
        from docx.oxml import OxmlElement
        doc=Document();doc.add_paragraph('Перед таблицей')
        t=doc.add_table(rows=3,cols=3);t.cell(0,0).text='Оборудование'
        t.cell(0,1).text='Количество';t.cell(0,1).merge(t.cell(0,2))
        t.cell(1,0).text='Насос';t.cell(1,0).merge(t.cell(2,0));t.cell(1,1).text='0'
        nested=t.cell(2,1).add_table(rows=1,cols=1);nested.cell(0,0).text='Вложенное'
        flag=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(flag)
        doc.add_paragraph('После таблицы')
        h=doc.sections[0].header.add_table(rows=1,cols=1,width=1000000);h.cell(0,0).text='Колонтитул'
        doc.save(self.sources/'data.docx')
        store=self.prepare();store.build();sheets=store.catalog()['items']
        body=next(s for s in sheets if s['name']=='word/document.xml')
        self.assertFalse(body['physical_page']);self.assertEqual(len(body['table_ids']),2)
        self.assertEqual([t['table_id'] for t in store.catalog('tables')['items'] if t['unit_id']==body['unit_id']],body['table_ids'])
        self.assertEqual(body['blocks'][0]['kind'],'p');self.assertEqual(body['blocks'][1]['kind'],'tbl')
        main=store.read(body['table_ids'][0])
        self.assertEqual(main['table']['header_rows'],[1]);self.assertEqual(main['rows'][0][1]['column_span'],2)
        self.assertEqual(main['rows'][1][0]['row_span'],2)
        self.assertEqual(main['rows'][2][0]['state'],'merged_continuation')
        # The stored address must resolve in the original mixed paragraph/table XML order.
        import zipfile,xml.etree.ElementTree as ET
        with zipfile.ZipFile(self.sources/'data.docx') as z:root=ET.fromstring(z.read('word/document.xml'))
        locator=main['rows'][1][0]['locator']['xml_path']
        xpath='/'.join('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'+segment for segment in locator.split('/')[2:])
        self.assertIsNotNone(root.find(xpath),locator)
        self.assertNotIn('Вложенное',json.dumps(main['rows'],ensure_ascii=False))
        child=store.read(body['table_ids'][1]);self.assertEqual(child['table']['parent_tables'],[1])
        self.assertEqual(child['rows'][0][0]['text'],'Вложенное')
        self.assertEqual(len(sheets),2)

    def test_pdf_physical_pages_rotation_scan_and_cache_without_reparse(self):
        import pymupdf
        from test_document_tables import grid
        from unittest.mock import patch
        with pymupdf.open() as pdf:
            p=pdf.new_page(width=500,height=400)
            p.insert_font(fontname='fixture',fontbuffer=pymupdf.Font('cjk').buffer)
            grid(p,30,30,[['Марка','Кол.'],['П1','0']],[70,60]);grid(p,250,30,[['Марка','Кол.'],['П1','0']],[70,60])
            p.set_cropbox(pymupdf.Rect(10,10,490,390));p.set_rotation(90);pdf.new_page()
            pdf.save(self.sources/'data.pdf')
        store=self.prepare();store.build(unit_limit=1)
        self.assertEqual(store.coverage()['pending_documents'],1)
        store.build(unit_limit=1)
        self.assertEqual(store.coverage()['partial_documents'],1)
        pages=store.catalog()['items'];self.assertTrue(pages[0]['physical_page'])
        self.assertEqual(pages[0]['original_rotation'],90);self.assertEqual(len(pages[0]['table_ids']),2)
        self.assertIn('OCR',pages[1]['warnings'][0])
        with patch('uvyazka.document_tables.detect',side_effect=AssertionError('must read cache')):
            table=store.read(pages[0]['table_ids'][0])
        self.assertEqual(table['rows'][1][1]['text'],'0')
        self.assertIsNotNone(table['rows'][1][1]['locator']['bbox'])
        self.assertEqual(table['rows'][1][1]['locator']['coordinate_system'],'pdf_unrotated_top_left_points')

    def test_two_reading_lanes_history_idempotence_stale_and_project_boundary(self):
        source=self.sources/'data.csv';source.write_text('Площадь,12.50\nПусто,\n',encoding='utf-8')
        store=self.prepare();store.build();tid=store.catalog('tables')['items'][0]['table_id']
        c=store.read(tid)['rows'][0][1];cid=c['cell_id']
        reading=store.review(cid,'model','12,5','value','visual_original','one')
        self.assertEqual(reading['comparison'],'matches_numeric');self.assertFalse(reading['accepted'])
        self.assertEqual(store.review(cid,'model','12,5','value','visual_original','one'),reading)
        with self.assertRaisesRegex(ValueError,'request_id'):store.review(cid,'model','9','value','visual_original','one')
        with self.assertRaisesRegex(ValueError,'conflict'):store.review(cid,'model','9','value','visual_original','two')
        human=store.review(cid,'human','10','value','cached_extraction','human')
        self.assertEqual(human['comparison'],'discrepancy');self.assertEqual(human['independence'],'not_independent')
        store.review(cid,'model','12.50','value','native_source','two','Повторное чтение',1)
        self.assertEqual(store.read(tid)['rows'][0][1]['readings']['model']['revision'],2)
        with connection(store.path,readonly=True) as db:self.assertEqual(db.execute('select count(*) from structure_reviews').fetchone()[0],3)
        source.write_text('Площадь,99\n',encoding='utf-8')
        self.assertFalse(store.read(tid)['current_registered_version'])
        stale=store.review(cid,'human','99','value','visual_original','stale','Изменился файл',1)
        self.assertEqual(stale['comparison'],'stale')
        scan(self.workspace,'a');self.assertEqual(store.catalog('tables')['total'],0)
        store.build();self.assertEqual(store.catalog('tables')['total'],1)
        self.assertFalse(store.read(tid)['current_registered_version'])
        self.workspace.create_project('b','Other',self.sources)
        with self.assertRaisesRegex(ValueError,'Unknown table'):ProjectStructure(self.workspace,'b').read(tid)

    def test_row_column_windows_and_unsupported_coverage(self):
        (self.sources/'grid.csv').write_text('\n'.join(','.join(str(r*30+c) for c in range(30)) for r in range(503)),encoding='utf-8')
        (self.sources/'old.doc').write_bytes(b'unsupported');(self.sources/'old.odt').write_bytes(b'unsupported')
        store=self.prepare();store.build();self.assertEqual(store.coverage()['unsupported_documents'],2)
        tid=store.catalog('tables')['items'][0]['table_id'];data=store.read(tid,offset=500,column_offset=25)
        self.assertEqual(len(data['rows']),3);self.assertEqual(len(data['rows'][0]),5)
        self.assertEqual(data['rows'][0][0]['text'],'15025')
        self.assertIsNone(data['next_offset']);self.assertIsNone(data['next_column_offset'])
        with self.assertRaises(ValueError):store.read(tid,limit=1000)

    def test_legacy_xls_keeps_formula_limit_and_native_merge(self):
        import shutil
        shutil.copyfile(Path(__file__).parent/'fixtures/structure-synthetic.xls',self.sources/'data.xls')
        store=self.prepare();store.build()
        self.assertEqual(store.coverage()['partial_documents'],1)
        self.assertEqual(store.coverage()['unsupported_documents'],0)
        sheet=store.catalog()['items'][0];data=store.read(sheet['table_ids'][0])
        self.assertEqual(data['rows'][0][1]['state'],'merged_continuation')
        self.assertEqual(data['rows'][1][1]['text'],'0.0')
        self.assertEqual(data['rows'][2][1]['formula_status'],'unknown_not_exposed_by_xlrd')
        self.assertIn('formula_expressions_unavailable',sheet['warnings'][0])
        doc=self.workspace.find_documents('a')['documents'][0]
        self.assertEqual(doc['index_status'],'text_indexed')

    def test_failure_is_explicit_and_retry_makes_progress_once(self):
        from unittest.mock import patch
        (self.sources/'data.csv').write_text('П1,10\n',encoding='utf-8');store=self.prepare()
        with patch('uvyazka.project_structure.extract',side_effect=ValueError('temporary reader failure')):
            self.assertEqual(store.build()['coverage']['failed_documents'],1)
        self.assertEqual(store.build()['processed'],[])
        self.assertEqual(store.build(retry_failed=True)['coverage']['indexed_documents'],1)

    def test_folder_intake_builds_structure_for_both_clients(self):
        from uvyazka.intake import ProjectIntake
        (self.sources/'data.csv').write_text('Нагрузка,10,kW\n',encoding='utf-8')
        intake=ProjectIntake(self.workspace)
        job=intake.import_folder(self.sources,background=False)
        result=intake.run(job['intake_id'])
        self.assertNotEqual(result['status'],'failed',result)
        service=ProjectStructure(Workspace(self.workspace.home),result['project_id'])
        self.assertEqual(service.coverage()['tables'],1)
        self.assertIn('get_project_structure',result['next_tools'])

    def test_http_same_db_and_write_protection(self):
        from uvyazka.dashboard import make_handler
        (self.sources/'data.csv').write_text('П1,10\n',encoding='utf-8');self.prepare()
        server=ThreadingHTTPServer(('127.0.0.1',0),make_handler(self.workspace))
        thread=Thread(target=server.serve_forever,daemon=True);thread.start()
        base=f'http://127.0.0.1:{server.server_port}/api/structure/'
        def get(path):
            with urlopen(base+path) as response:return json.load(response)
        try:
            token=get('config')['write_token']
            req=Request(base+'build?project_id=a',data=b'{}',headers={'Content-Type':'application/json'},method='POST')
            with self.assertRaises(HTTPError) as error:urlopen(req)
            self.assertEqual(error.exception.code,403)
            req.add_header('X-Uvyazka-Token',token)
            with urlopen(req) as response:self.assertEqual(json.load(response)['coverage']['tables'],1)
            tid=get('catalog?project_id=a&kind=tables')['items'][0]['table_id']
            c=get('table?project_id=a&table_id='+tid)['rows'][0][1]
            payload=dict(cell_id=c['cell_id'],reader='human',observed_text='10',state='value',method='cached_extraction',request_id='http')
            req=Request(base+'review?project_id=a',data=json.dumps(payload).encode(),headers={'Content-Type':'application/json','X-Uvyazka-Token':token},method='POST')
            with urlopen(req) as response:self.assertEqual(json.load(response)['comparison'],'matches_literal')
            self.assertIn('human',ProjectStructure(self.workspace,'a').read(tid)['rows'][0][1]['readings'])
            payload=dict(analysis_id='http-note',table_id=tid,title='Количество',summary='Вторая колонка — число.',data={'columns':{'2':'quantity'}},reader='human',request_id='http-analysis')
            req=Request(base+'analysis?project_id=a',data=json.dumps(payload).encode(),headers={'Content-Type':'application/json','X-Uvyazka-Token':token},method='POST')
            with urlopen(req) as response:self.assertEqual(json.load(response)['revision'],1)
            self.assertEqual(ProjectStructure(self.workspace,'a').read(tid)['analyses']['items'][0]['analysis_id'],'http-note')
            req.add_header('Origin','https://example.org')
            with self.assertRaises(HTTPError) as error:urlopen(req)
            self.assertEqual(error.exception.code,403)
        finally:server.shutdown();server.server_close();thread.join()

    def test_analysis_links_two_sources_keeps_lanes_history_and_staleness(self):
        from uvyazka.structure_analysis import StructureAnalysis
        (self.sources/'spec.csv').write_text('Лоток,1161\n',encoding='utf-8')
        other=self.sources/'work.csv';other.write_text('Подвес,991\nСтена,170\n',encoding='utf-8')
        store=self.prepare();store.build();tables=store.catalog('tables')['items'];a,b=tables
        notes=StructureAnalysis(self.workspace,'a')
        args=dict(analysis_id='tray',table_id=a['table_id'],title='Лоток и монтаж',summary='Количество работ нужно складывать по способу крепления.',
                  data={'columns':{'2':'quantity'},'mapping':{'operation':'sum','expected':'1161'},'questions':['Область должна совпадать']},
                  reader='model',request_id='analysis-one',
                  related_tables=[dict(table_id=b['table_id'],relation='Монтаж',reason='Один материал, два способа крепления')],
                  evidence_cells=[a['table_id']+':1:2',b['table_id']+':1:2',b['table_id']+':2:2'])
        saved=notes.save(**args);self.assertFalse(saved['accepted']);self.assertEqual(notes.save(**args),saved)
        args['data']['mapping']['expected']='99'
        with self.assertRaisesRegex(ValueError,'request_id'):notes.save(**args)
        self.assertEqual(store.read(a['table_id'])['analyses']['items'][0]['data']['mapping']['expected'],'1161')
        args.update(request_id='analysis-two',expected_revision=1,reason='Уточнение формулировки',summary='Нужно 991 + 170.')
        notes.save(**args)
        args.update(reader='human',expected_revision=0,reason='',request_id='analysis-human')
        notes.save(**args)
        latest=store.catalog('analyses');self.assertEqual(latest['total'],2)
        self.assertEqual(store.catalog('analyses',history=True)['total'],3)
        self.assertEqual(notes.list(table_id=b['table_id'])['total'],2)
        other.write_text('Подвес,992\nСтена,170\n',encoding='utf-8')
        items=store.read(a['table_id'])['analyses']['items'];self.assertTrue(all(not item['current_sources'] for item in items))
        args.update(request_id='analysis-stale',expected_revision=1,reason='Попытка после изменения файла')
        with self.assertRaisesRegex(ValueError,'Source changed'):notes.save(**args)

    def test_analysis_rejects_unknown_cells_cross_project_and_nonfinite_json(self):
        from uvyazka.structure_analysis import StructureAnalysis
        (self.sources/'table.csv').write_text('Насос,10\n',encoding='utf-8');store=self.prepare();store.build()
        tid=store.catalog('tables')['items'][0]['table_id'];notes=StructureAnalysis(self.workspace,'a')
        args=dict(analysis_id='a',table_id=tid,title='Разбор',summary='Пояснение',data={},reader='model',request_id='request')
        with self.assertRaisesRegex(ValueError,'Unknown evidence'):notes.save(**args,evidence_cells=[tid+':100:100'])
        with self.assertRaisesRegex(ValueError,'finite JSON'):notes.save(**{**args,'data':{'bad':float('nan')}})
        self.workspace.create_project('b','Other',self.sources)
        with self.assertRaisesRegex(ValueError,'Build source structure'):StructureAnalysis(self.workspace,'b').save(**args)

    def test_scoped_scan_leaves_unselected_sources_and_old_evidence_intact(self):
        (self.sources/'a.txt').write_text('one 1 m2',encoding='utf-8')
        other=self.sources/'b.txt';other.write_text('two 2 m2',encoding='utf-8');self.prepare()
        docs=self.workspace.find_documents('a')['documents'];a,b=docs
        (self.sources/'a.txt').write_text('one 3 m2',encoding='utf-8');other.unlink()
        scan(self.workspace,'a',doc_ids=[a['doc_id']])
        with connection(self.store.path,readonly=True) as db:
            row=dict(db.execute('select * from documents where doc_id=?',(b['doc_id'],)).fetchone())
        self.assertEqual(row['index_status'],'text_indexed')
        self.assertIn('one 1',self.workspace.read_evidence('a',a['doc_id'],sha256=a['sha256'])['text'])
        with self.assertRaises(ValueError):scan(self.workspace,'a',doc_ids=['not-registered'])
        from uvyazka.project_summary import ProjectSummary
        summary=ProjectSummary(self.workspace,'a')
        result=summary.index(doc_ids=[a['doc_id']])
        self.assertEqual(result['processed_pages'],1);self.assertEqual(result['scope_remaining_blocks'],0)
        self.assertEqual(result['remaining_blocks'],1)
        with connection(self.store.path,readonly=True) as db:
            self.assertEqual(db.execute('select count(*) from quantity_pages where doc_id=?',(b['doc_id'],)).fetchone()[0],0)

    def test_selected_pdf_pages_do_not_imply_preceding_pages_were_read(self):
        import pymupdf
        with pymupdf.open() as pdf:
            for i in range(5):pdf.new_page().insert_text((40,40),f'Synthetic page {i+1}')
            pdf.save(self.sources/'pages.pdf')
        store=self.prepare();doc=self.workspace.find_documents('a')['documents'][0]['doc_id']
        first=store.build(doc,unit_limit=1,start_unit=5)
        self.assertEqual(first['coverage']['indexed_units'],1);self.assertEqual(first['processed'][0]['next_unit'],0)
        self.assertEqual(first['processed'][0]['next_unread_unit'],1)
        self.assertEqual(first['scope_coverage']['documents'],1)
        self.assertEqual(store.catalog()['items'][0]['index'],5)
        store.build(doc,unit_limit=2)
        self.assertEqual(store.coverage()['indexed_units'],3)
        # A targeted range may overlap already indexed units and extend beyond the first gap.
        store.build(doc,unit_limit=3,start_unit=2)
        self.assertEqual(store.coverage()['indexed_units'],5);self.assertEqual(store.coverage()['pending_documents'],0)
        with self.assertRaisesRegex(ValueError,'outside source'):store.build(doc,start_unit=6)


if __name__=='__main__':unittest.main()
